"""Imitation-learning module for real-world train traction strategies.

This module is intentionally independent from the existing DQN lower controller.
It learns a continuous force-ratio policy from measured operation records and can
later be used as a warm-start policy or as a smooth action prior.
"""

from __future__ import annotations

import csv
import math
import random
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
from openpyxl import load_workbook
import torch
import torch.nn as nn
import torch.optim as optim
from torch.utils.data import DataLoader, Dataset, random_split


TIME_ALIASES = (
    "time_s",
    "timestamp_s",
    "elapsed_s",
    "time",
    "timestamp",
    "\u65f6\u95f4",
    "\u65f6\u95f4\u6233",
)
POSITION_ALIASES = (
    "\u65f6\u95f4\u4f4d\u7f6e",
    "position_m",
    "mileage_m",
    "distance_m",
    "position",
    "mileage",
    "\u4f4d\u7f6e",
    "\u91cc\u7a0b",
)
SPEED_ALIASES = (
    "\u5217\u8f66\u7efc\u5408\u901f\u5ea6\uff081=0.1km/h\uff09",
    "speed_mps",
    "speed_kmh",
    "speed",
    "velocity",
    "\u901f\u5ea6",
)
FORCE_ALIASES = (
    "\u5408\u529b",
    "force_n",
    "traction_force_n",
    "motor_force_n",
    "force_kn",
    "traction_force_kn",
    "force",
    "traction_force",
    "\u7275\u5f15\u529b",
    "\u7535\u673a\u51fa\u529b",
    "\u7275\u5f15\u5236\u52a8\u529b",
)
SPEED_LIMIT_ALIASES = (
    "speed_limit_kmh",
    "limit_kmh",
    "speed_limit",
    "\u9650\u901f",
    "\u6700\u5927\u9650\u901f",
)
GRADIENT_ALIASES = (
    "gradient_per_mille",
    "gradient",
    "slope",
    "\u5761\u5ea6",
)
MOTOR_CAR_MASS_ALIASES = (
    "\u52a8\u8f66\u8d28\u91cf",
    "motor_car_mass",
    "motor_mass",
)
TRAILER_CAR_MASS_ALIASES = (
    "\u62d6\u8f66\u8d28\u91cf",
    "trailer_car_mass",
    "trailer_mass",
)
GROUP_ALIASES = (
    "run_id",
    "trip_id",
    "train_id",
    "section_name",
    "\u7ad9\u95f4",
    "\u8f66\u6b21",
)


@dataclass(frozen=True)
class RealStrategyConfig:
    """Training knobs for the real-traction behavior cloning policy."""

    history_len: int = 12
    max_abs_force_n: float = 300000.0
    speed_scale_mps: float = 40.0
    acceleration_scale_mps2: float = 1.5
    position_scale_m: float = 6000.0
    gradient_scale_per_mille: float = 35.0
    mass_scale_kg: float = 300000.0
    max_abs_gradient_per_mille: float = 200.0
    max_speed_kmh: float = 130.0
    max_abs_force_n_for_filter: float = 400000.0
    min_train_mass_kg: float = 180000.0
    max_train_mass_kg: float = 360000.0
    max_position_backtrack_m: float = 5.0
    force_smoothness_weight: float = 0.08
    sample_interval_s: Optional[float] = 0.1
    execution_tracking_weight: float = 0.35
    execution_tracking_near_weight: float = 0.75
    execution_tracking_response_time_s: float = 8.0
    hidden_dim: int = 96
    gru_layers: int = 1
    dropout: float = 0.0


@dataclass
class RealStrategyStats:
    sample_count: int
    accepted_sample_count: int
    filtered_sample_count: int
    filter_reasons: Dict[str, int]
    sequence_count: int
    feature_names: List[str]
    train_loss: float
    validation_loss: float


def _canonical(name: str) -> str:
    return (
        name.strip()
        .lower()
        .replace("\uff08", "(")
        .replace("\uff09", ")")
        .replace(" ", "_")
    )


def _find_column(columns: Sequence[str], aliases: Sequence[str]) -> Optional[str]:
    canonical_to_raw = {_canonical(column): column for column in columns}
    for alias in aliases:
        key = _canonical(alias)
        if key in canonical_to_raw:
            return canonical_to_raw[key]
    for column in columns:
        key = _canonical(column)
        if any(_canonical(alias) in key for alias in aliases):
            return column
    return None


def _to_float(value: object) -> Optional[float]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _section_name_from_path(path: Path) -> str:
    stem = path.stem
    if "-" not in stem:
        return stem
    prefix, section_name = stem.split("-", 1)
    return section_name if prefix.isdigit() and section_name else stem


def _expand_paths(paths: Sequence[Path]) -> List[Path]:
    files: List[Path] = []
    for path in paths:
        if path.is_dir():
            files.extend(sorted(path.glob("*.csv")))
            files.extend(sorted(path.glob("*.xlsx")))
        else:
            files.append(path)
    return files


def _read_rows(paths: Sequence[Path]) -> List[Dict[str, object]]:
    rows: List[Dict[str, object]] = []
    for path in _expand_paths(paths):
        section_name = _section_name_from_path(path)
        suffix = path.suffix.lower()
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                for idx, row in enumerate(reader):
                    row = dict(row)
                    row.setdefault("section_name", section_name)
                    row["_source_file"] = str(path)
                    row["_source_row"] = idx
                    row["_sample_index"] = idx
                    rows.append(row)
        elif suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                iterator = worksheet.iter_rows(values_only=True)
                try:
                    headers = list(next(iterator))
                except StopIteration:
                    continue
                if not any(header is not None for header in headers):
                    continue
                columns = [
                    str(header) if header is not None else f"unnamed_{idx + 1}"
                    for idx, header in enumerate(headers)
                ]
                for idx, values in enumerate(iterator):
                    if values is None or all(value is None for value in values):
                        continue
                    row = {
                        column: values[col_idx] if col_idx < len(values) else None
                        for col_idx, column in enumerate(columns)
                    }
                    row["section_name"] = section_name
                    row["_source_file"] = str(path)
                    row["_source_sheet"] = sheet_name
                    row["_source_row"] = idx
                    row["_sample_index"] = idx
                    rows.append(row)
        else:
            raise ValueError(f"Unsupported real-operation data file: {path}")
    return rows


def _infer_speed_mps(value: float, column_name: str) -> float:
    key = _canonical(column_name)
    if "1=0.1km/h" in key or "0.1km/h" in key:
        return value * 0.1 / 3.6
    if "kmh" in key or "km/h" in key or "\u516c\u91cc" in key:
        return value / 3.6
    if abs(value) > 45.0:
        return value / 3.6
    return value


def _infer_force_n(value: float, column_name: str, global_force_abs_max: float) -> float:
    key = _canonical(column_name)
    if "kn" in key or "\u5343\u725b" in key:
        return value * 1000.0
    if global_force_abs_max < 2000.0:
        return value * 1000.0
    return value


def _group_key(row: Dict[str, object], group_columns: Sequence[str]) -> Tuple[str, ...]:
    if group_columns:
        return tuple(str(row.get(column, "")) for column in group_columns)
    values = [str(row.get("_source_file", "")), str(row.get("_source_sheet", ""))]
    for alias in GROUP_ALIASES:
        column = next((key for key in row if _canonical(key) == _canonical(alias)), None)
        if column is not None:
            values.append(str(row.get(column, "")))
    return tuple(values)


class RealOperationDataset(Dataset):
    """Fixed-history windows from measured train-operation CSV records."""

    feature_names = [
        "speed_mps",
        "position_progress",
        "time_progress",
        "acceleration_mps2",
        "previous_force_ratio",
        "speed_limit_kmh",
        "gradient_per_mille",
        "train_mass_kg",
    ]

    def __init__(
        self,
        data_paths: Sequence[Path],
        config: RealStrategyConfig,
        group_columns: Sequence[str] = (),
    ) -> None:
        rows = _read_rows(data_paths)
        if not rows:
            raise ValueError("No CSV rows were found in the provided real-operation data")

        columns = sorted({column for row in rows for column in row.keys()})
        time_column = _find_column(columns, TIME_ALIASES)
        position_column = _find_column(columns, POSITION_ALIASES)
        fallback_position_column = _find_column(
            columns,
            (
                "position_m",
                "mileage_m",
                "distance_m",
                "position",
                "mileage",
                "\u4f4d\u7f6e",
                "\u91cc\u7a0b",
            ),
        )
        speed_column = _find_column(columns, SPEED_ALIASES)
        force_column = _find_column(columns, FORCE_ALIASES)
        speed_limit_column = _find_column(columns, SPEED_LIMIT_ALIASES)
        gradient_column = _find_column(columns, GRADIENT_ALIASES)
        motor_mass_column = _find_column(columns, MOTOR_CAR_MASS_ALIASES)
        trailer_mass_column = _find_column(columns, TRAILER_CAR_MASS_ALIASES)
        if position_column is None or speed_column is None or force_column is None:
            raise ValueError(
                "Real-operation data must contain position, speed and force columns. "
                f"Detected columns: {columns}"
            )

        raw_force_values = [
            abs(value)
            for value in (_to_float(row.get(force_column)) for row in rows)
            if value is not None
        ]
        global_force_abs_max = max(raw_force_values) if raw_force_values else config.max_abs_force_n

        grouped: Dict[Tuple[str, ...], List[Dict[str, float]]] = {}
        self.raw_sample_count = 0
        self.accepted_sample_count = 0
        self.filtered_sample_count = 0
        self.filter_reasons: Dict[str, int] = {}

        def reject(reason: str) -> None:
            self.filtered_sample_count += 1
            self.filter_reasons[reason] = self.filter_reasons.get(reason, 0) + 1

        for row in rows:
            self.raw_sample_count += 1
            if config.sample_interval_s is not None:
                sample_index = _to_float(row.get("_sample_index"))
                time_s = None if sample_index is None else sample_index * config.sample_interval_s
            else:
                time_s = _to_float(row.get(time_column)) if time_column else None
            position_m = _to_float(row.get(position_column))
            if position_m is None and fallback_position_column is not None:
                position_m = _to_float(row.get(fallback_position_column))
            speed_value = _to_float(row.get(speed_column))
            force_value = _to_float(row.get(force_column))
            if None in (time_s, position_m, speed_value, force_value):
                reject("missing_required_value")
                continue
            speed_mps = _infer_speed_mps(float(speed_value), speed_column)
            force_n = _infer_force_n(float(force_value), force_column, global_force_abs_max)
            speed_limit = _to_float(row.get(speed_limit_column)) if speed_limit_column else None
            gradient = _to_float(row.get(gradient_column)) if gradient_column else None
            motor_mass = _to_float(row.get(motor_mass_column)) if motor_mass_column else None
            trailer_mass = _to_float(row.get(trailer_mass_column)) if trailer_mass_column else None
            train_mass_kg = 0.0
            if motor_mass is not None or trailer_mass is not None:
                # The source columns are already aggregate motor/trailer mass.
                # Values in these files are ton-scale, so convert to kg.
                train_mass_kg = (float(motor_mass or 0.0) + float(trailer_mass or 0.0)) * 1000.0
            gradient_value = float(gradient) if gradient is not None else 0.0
            if speed_mps < -1e-6 or speed_mps * 3.6 > config.max_speed_kmh:
                reject("speed_out_of_range")
                continue
            if abs(force_n) > config.max_abs_force_n_for_filter:
                reject("force_out_of_range")
                continue
            if abs(gradient_value) > config.max_abs_gradient_per_mille:
                reject("gradient_out_of_range")
                continue
            if train_mass_kg > 0.0 and not (
                config.min_train_mass_kg <= train_mass_kg <= config.max_train_mass_kg
            ):
                reject("mass_out_of_range")
                continue
            self.accepted_sample_count += 1
            grouped.setdefault(_group_key(row, group_columns), []).append(
                {
                    "time_s": float(time_s),
                    "position_m": float(position_m),
                    "speed_mps": float(speed_mps),
                    "force_n": float(force_n),
                    "speed_limit_kmh": float(speed_limit) if speed_limit is not None else 0.0,
                    "gradient_per_mille": gradient_value,
                    "train_mass_kg": train_mass_kg,
                }
            )

        self.windows: List[np.ndarray] = []
        self.targets: List[float] = []
        self.previous_force_ratios: List[float] = []
        self.config = config
        history_len = max(2, int(config.history_len))
        for group_rows in grouped.values():
            group_rows = sorted(group_rows, key=lambda item: (item["time_s"], item["position_m"]))
            if len(group_rows) < history_len:
                continue
            cleaned_rows: List[Dict[str, float]] = []
            previous_position: Optional[float] = None
            for item in group_rows:
                if (
                    previous_position is not None
                    and item["position_m"] < previous_position - config.max_position_backtrack_m
                ):
                    self.filtered_sample_count += 1
                    self.filter_reasons["position_backtrack"] = (
                        self.filter_reasons.get("position_backtrack", 0) + 1
                    )
                    continue
                cleaned_rows.append(item)
                previous_position = item["position_m"]
            group_rows = cleaned_rows
            if len(group_rows) < history_len:
                continue
            start_time = group_rows[0]["time_s"]
            end_time = max(group_rows[-1]["time_s"], start_time + 1.0)
            start_position = group_rows[0]["position_m"]
            end_position = group_rows[-1]["position_m"]
            position_span = max(abs(end_position - start_position), config.position_scale_m, 1.0)
            previous_speed = group_rows[0]["speed_mps"]
            previous_time = group_rows[0]["time_s"]
            previous_force_ratio = 0.0
            features: List[List[float]] = []
            force_ratios: List[float] = []
            for item in group_rows:
                dt_s = max(item["time_s"] - previous_time, 1e-3)
                acceleration = (item["speed_mps"] - previous_speed) / dt_s
                force_ratio = float(np.clip(item["force_n"] / config.max_abs_force_n, -1.0, 1.0))
                features.append(
                    [
                        item["speed_mps"] / config.speed_scale_mps,
                        (item["position_m"] - start_position) / position_span,
                        (item["time_s"] - start_time) / max(end_time - start_time, 1.0),
                        acceleration / config.acceleration_scale_mps2,
                        previous_force_ratio,
                        item["speed_limit_kmh"] / 125.0,
                        item["gradient_per_mille"] / config.gradient_scale_per_mille,
                        item["train_mass_kg"] / config.mass_scale_kg,
                    ]
                )
                force_ratios.append(force_ratio)
                previous_speed = item["speed_mps"]
                previous_time = item["time_s"]
                previous_force_ratio = force_ratio

            feature_array = np.asarray(features, dtype=np.float32)
            target_array = np.asarray(force_ratios, dtype=np.float32)
            for end_idx in range(history_len - 1, len(group_rows)):
                start_idx = end_idx - history_len + 1
                self.windows.append(feature_array[start_idx : end_idx + 1])
                self.targets.append(float(target_array[end_idx]))
                self.previous_force_ratios.append(float(feature_array[end_idx, 4]))

        if not self.windows:
            raise ValueError("No trainable history windows were produced from the real-operation data")

    def __len__(self) -> int:
        return len(self.windows)

    def __getitem__(self, index: int) -> Tuple[torch.Tensor, torch.Tensor, torch.Tensor]:
        return (
            torch.as_tensor(self.windows[index], dtype=torch.float32),
            torch.as_tensor([self.targets[index]], dtype=torch.float32),
            torch.as_tensor([self.previous_force_ratios[index]], dtype=torch.float32),
        )


class RealTractionPolicyNet(nn.Module):
    """Continuous traction/braking force-ratio policy learned from real data."""

    def __init__(self, feature_dim: int, config: RealStrategyConfig) -> None:
        super().__init__()
        self.encoder = nn.GRU(
            input_size=feature_dim,
            hidden_size=config.hidden_dim,
            num_layers=config.gru_layers,
            batch_first=True,
            dropout=config.dropout if config.gru_layers > 1 else 0.0,
        )
        self.head = nn.Sequential(
            nn.Linear(config.hidden_dim, config.hidden_dim),
            nn.Tanh(),
            nn.Linear(config.hidden_dim, 1),
            nn.Tanh(),
        )

    def forward(self, features: torch.Tensor) -> torch.Tensor:
        encoded, _ = self.encoder(features)
        return self.head(encoded[:, -1, :])


class RealStrategyController:
    """Adapter that lets the real-data policy drive LowerSectionEnvironment."""

    def __init__(
        self,
        model: RealTractionPolicyNet,
        config: RealStrategyConfig,
        device_name: Optional[str] = None,
    ) -> None:
        self.config = config
        self.device = torch.device(device_name or ("cuda" if torch.cuda.is_available() else "cpu"))
        self.model = model.to(self.device)
        self.model.eval()
        self.histories: Dict[int, List[np.ndarray]] = {}
        self.previous_speed_mps: Dict[int, float] = {}
        self.previous_time_s: Dict[int, float] = {}

    @classmethod
    def from_checkpoint(
        cls,
        checkpoint_path: Path,
        device_name: Optional[str] = None,
    ) -> "RealStrategyController":
        checkpoint = torch.load(checkpoint_path, map_location=device_name or "cpu")
        config = RealStrategyConfig(**checkpoint["config"])
        model = RealTractionPolicyNet(len(checkpoint["feature_names"]), config)
        model.load_state_dict(checkpoint["model_state_dict"])
        return cls(model, config, device_name=device_name)

    def reset_environment(self, environment_id: int) -> None:
        self.histories.pop(environment_id, None)
        self.previous_speed_mps.pop(environment_id, None)
        self.previous_time_s.pop(environment_id, None)

    def _feature_from_environment(self, environment) -> np.ndarray:
        key = id(environment)
        if environment.step_count == 0:
            self.reset_environment(key)
        segment = environment.section.segment_at(environment.position_m)
        previous_speed = self.previous_speed_mps.get(key, environment.speed_mps)
        previous_time = self.previous_time_s.get(key, environment.elapsed_s)
        dt_s = max(environment.elapsed_s - previous_time, 1e-3)
        acceleration = (environment.speed_mps - previous_speed) / dt_s
        position_progress = (
            (environment.position_m - environment.section.start_m)
            / max(environment.section.length_m, 1.0)
        )
        time_progress = environment.elapsed_s / max(environment.target_time_s, 1.0)
        previous_force_ratio = float(
            np.clip(environment.previous_force_n / max(self.config.max_abs_force_n, 1.0), -1.0, 1.0)
        )
        self.previous_speed_mps[key] = environment.speed_mps
        self.previous_time_s[key] = environment.elapsed_s
        return np.asarray(
            [
                environment.speed_mps / self.config.speed_scale_mps,
                position_progress,
                time_progress,
                acceleration / self.config.acceleration_scale_mps2,
                previous_force_ratio,
                segment.speed_limit_kmh / 125.0,
                segment.gradient_per_mille / self.config.gradient_scale_per_mille,
                environment.dynamics.config.mass_kg / self.config.mass_scale_kg,
            ],
            dtype=np.float32,
        )

    def select_action_and_force(self, environment) -> Tuple[int, float]:
        key = id(environment)
        feature = self._feature_from_environment(environment)
        history = self.histories.setdefault(key, [])
        history.append(feature)
        history_len = max(2, int(self.config.history_len))
        if len(history) > history_len:
            del history[:-history_len]
        if len(history) < history_len:
            padded = [history[0]] * (history_len - len(history)) + history
        else:
            padded = history
        window = torch.as_tensor(
            np.asarray(padded, dtype=np.float32),
            dtype=torch.float32,
            device=self.device,
        ).unsqueeze(0)
        with torch.no_grad():
            force_ratio = float(self.model(window).squeeze().item())
        continuous_force_n = force_ratio * self.config.max_abs_force_n
        remaining_time_s = environment.target_time_s - environment.elapsed_s
        if remaining_time_s > 1.0:
            required_speed_mps = environment.remaining_m / max(remaining_time_s, 1.0)
            speed_error_mps = required_speed_mps - environment.speed_mps
            desired_acceleration_mps2 = np.clip(
                speed_error_mps / max(self.config.execution_tracking_response_time_s, 1.0),
                -0.65,
                0.65,
            )
            segment = environment.section.segment_at(environment.position_m)
            resistance_n = environment.dynamics.total_resistance_n(environment.speed_mps, segment)
            tracking_force_n = resistance_n + environment.dynamics.config.mass_kg * desired_acceleration_mps2
            tracking_force_n = float(
                np.clip(
                    tracking_force_n,
                    -environment.dynamics.config.max_braking_force_n * 0.65,
                    environment.dynamics.max_traction_force_n(environment.speed_mps),
                )
            )
            near_terminal = environment.remaining_m < 1600.0 or remaining_time_s < 100.0
            tracking_weight = (
                self.config.execution_tracking_near_weight
                if near_terminal
                else self.config.execution_tracking_weight
            )
            continuous_force_n = (
                (1.0 - tracking_weight) * continuous_force_n
                + tracking_weight * tracking_force_n
            )
        action = environment.action_for_force_n(continuous_force_n)
        return action, continuous_force_n


def train_real_strategy_policy(
    data_paths: Sequence[Path],
    output_path: Path,
    config: RealStrategyConfig,
    group_columns: Sequence[str] = (),
    epochs: int = 60,
    batch_size: int = 256,
    learning_rate: float = 3e-4,
    validation_ratio: float = 0.15,
    seed: int = 7,
    device_name: Optional[str] = None,
) -> RealStrategyStats:
    """Train and save a continuous real-traction behavior-cloning policy."""

    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    dataset = RealOperationDataset(data_paths, config, group_columns)
    validation_size = max(1, int(len(dataset) * validation_ratio))
    train_size = max(1, len(dataset) - validation_size)
    if train_size + validation_size > len(dataset):
        validation_size = len(dataset) - train_size
    generator = torch.Generator().manual_seed(seed)
    train_dataset, validation_dataset = random_split(
        dataset,
        [train_size, validation_size],
        generator=generator,
    )
    device = torch.device(device_name or ("cuda" if torch.cuda.is_available() else "cpu"))
    model = RealTractionPolicyNet(len(dataset.feature_names), config).to(device)
    optimizer = optim.AdamW(model.parameters(), lr=learning_rate)
    train_loader = DataLoader(train_dataset, batch_size=batch_size, shuffle=True)
    validation_loader = DataLoader(validation_dataset, batch_size=batch_size, shuffle=False)

    def run_epoch(loader: DataLoader, train: bool) -> float:
        model.train(train)
        losses: List[float] = []
        with torch.set_grad_enabled(train):
            for features, targets, previous_force in loader:
                features = features.to(device)
                targets = targets.to(device)
                previous_force = previous_force.to(device)
                predictions = model(features)
                imitation_loss = nn.functional.mse_loss(predictions, targets)
                smoothness_loss = nn.functional.mse_loss(predictions, previous_force)
                loss = imitation_loss + config.force_smoothness_weight * smoothness_loss
                if train:
                    optimizer.zero_grad()
                    loss.backward()
                    nn.utils.clip_grad_norm_(model.parameters(), 1.0)
                    optimizer.step()
                losses.append(float(loss.item()))
        return float(np.mean(losses)) if losses else math.nan

    train_loss = validation_loss = math.nan
    for epoch in range(max(1, epochs)):
        train_loss = run_epoch(train_loader, train=True)
        validation_loss = run_epoch(validation_loader, train=False)
        if epoch == 0 or (epoch + 1) % max(1, epochs // 10) == 0:
            print(
                f"  real-policy epoch {epoch + 1:4d}/{epochs}: "
                f"train_loss={train_loss:.6f} validation_loss={validation_loss:.6f}"
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    stats = RealStrategyStats(
        sample_count=dataset.raw_sample_count,
        accepted_sample_count=dataset.accepted_sample_count,
        filtered_sample_count=dataset.filtered_sample_count,
        filter_reasons=dict(dataset.filter_reasons),
        sequence_count=len(dataset),
        feature_names=list(dataset.feature_names),
        train_loss=float(train_loss),
        validation_loss=float(validation_loss),
    )
    torch.save(
        {
            "model_state_dict": model.state_dict(),
            "config": asdict(config),
            "feature_names": list(dataset.feature_names),
            "stats": asdict(stats),
            "module": "real_strategy_behavior_cloning",
        },
        output_path,
    )
    return stats


def _grouped_lengths(data_paths: Sequence[Path], group_columns: Sequence[str]) -> Iterable[List[object]]:
    rows = _read_rows(data_paths)
    grouped: Dict[Tuple[str, ...], List[object]] = {}
    for row in rows:
        grouped.setdefault(_group_key(row, group_columns), []).append(row)
    return grouped.values()
