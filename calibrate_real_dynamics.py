"""Physics-consistent calibration for real THP-TD operation data.

The calibration keeps the measured speed, position and force close to the
original records while reducing kinematic and train-dynamics residuals:

    dx/dt ~= v
    m dv/dt ~= F - R_davis(v) - R_gradient(x) - R_curve(x)

For the Guangzhou Metro files, only rows with column "变" not equal to zero are
used as valid position anchors. Other position samples are allowed to move so
the trajectory remains continuous.
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import matplotlib
import numpy as np
import torch
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt

try:
    from .config import DEFAULT_LINE_DIR
    from .line_profile import LineProfile, StationSection
    from .train_model import TrainConfig, TrainDynamics
except ImportError:
    from config import DEFAULT_LINE_DIR
    from line_profile import LineProfile, StationSection
    from train_model import TrainConfig, TrainDynamics


SPEED_COLUMN = "\u5217\u8f66\u7efc\u5408\u901f\u5ea6\uff081=0.1km/h\uff09"
FORCE_COLUMN = "\u5408\u529b"
POSITION_COLUMN = "\u4f4d\u7f6e"
VALID_POSITION_COLUMN = "\u53d8"
TRAILER_MASS_COLUMN = "\u62d6\u8f66\u8d28\u91cf"
MOTOR_MASS_COLUMN = "\u52a8\u8f66\u8d28\u91cf"


@dataclass
class SheetCalibrationData:
    sheet_name: str
    time_s: np.ndarray
    speed_obs_mps: np.ndarray
    force_obs_n: np.ndarray
    position_obs_m: np.ndarray
    position_valid: np.ndarray
    mass_kg: float


@dataclass(frozen=True)
class CalibrationWeights:
    speed_sigma_mps: float = 0.70
    force_sigma_n: float = 35000.0
    position_sigma_m: float = 8.0
    kinematic_sigma_mps: float = 0.80
    dynamics_sigma_n: float = 55000.0
    force_smooth_sigma_n: float = 25000.0
    speed_smooth_sigma_mps: float = 0.80
    position_step_sigma_m: float = 8.0


def default_thp_td_path() -> Path:
    train_data_path = Path.home() / "Desktop" / "train_data" / "3-THP-TD.xlsx"
    if train_data_path.exists():
        return train_data_path
    return DEFAULT_LINE_DIR / "3-THP-TD.xlsx"


def find_column(headers: Sequence[object], name: str) -> Optional[int]:
    for idx, header in enumerate(headers):
        if header == name:
            return idx
    return None


def to_float(value: object) -> Optional[float]:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_sheet_data(
    workbook_path: Path,
    sheet_name: str,
    dt_s: float,
    fallback_mass_kg: float,
) -> SheetCalibrationData:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    headers = list(next(iterator))
    speed_idx = find_column(headers, SPEED_COLUMN)
    force_idx = find_column(headers, FORCE_COLUMN)
    position_idx = find_column(headers, POSITION_COLUMN)
    valid_idx = find_column(headers, VALID_POSITION_COLUMN)
    trailer_mass_idx = find_column(headers, TRAILER_MASS_COLUMN)
    motor_mass_idx = find_column(headers, MOTOR_MASS_COLUMN)
    required = {
        "speed": speed_idx,
        "force": force_idx,
        "position": position_idx,
        "valid_position": valid_idx,
    }
    missing = [name for name, idx in required.items() if idx is None]
    if missing:
        raise ValueError(f"{sheet_name} is missing required columns: {missing}")

    speed_mps: List[float] = []
    force_n: List[float] = []
    position_m: List[float] = []
    position_valid: List[bool] = []
    mass_values: List[float] = []
    for row in iterator:
        speed_raw = to_float(row[speed_idx])
        force_raw = to_float(row[force_idx])
        position_raw = to_float(row[position_idx])
        valid_raw = to_float(row[valid_idx])
        if speed_raw is None or force_raw is None or position_raw is None:
            continue
        speed_mps.append(speed_raw * 0.1 / 3.6)
        force_n.append(force_raw * 1000.0)
        position_m.append(position_raw)
        position_valid.append(valid_raw is not None and abs(valid_raw) > 1e-9)
        if trailer_mass_idx is not None and motor_mass_idx is not None:
            trailer_mass = to_float(row[trailer_mass_idx])
            motor_mass = to_float(row[motor_mass_idx])
            if trailer_mass is not None and motor_mass is not None:
                mass_values.append((trailer_mass + motor_mass) * 1000.0)

    if len(speed_mps) < 20:
        raise ValueError(f"{sheet_name} has too few valid samples")
    mass_kg = float(np.median(mass_values)) if mass_values else fallback_mass_kg
    sample_count = len(speed_mps)
    return SheetCalibrationData(
        sheet_name=sheet_name,
        time_s=np.arange(sample_count, dtype=float) * dt_s,
        speed_obs_mps=np.asarray(speed_mps, dtype=float),
        force_obs_n=np.asarray(force_n, dtype=float),
        position_obs_m=np.asarray(position_m, dtype=float),
        position_valid=np.asarray(position_valid, dtype=bool),
        mass_kg=mass_kg,
    )


def initial_position_from_speed(data: SheetCalibrationData, dt_s: float) -> np.ndarray:
    x = np.zeros_like(data.position_obs_m)
    valid_indices = np.flatnonzero(data.position_valid)
    x[0] = data.position_obs_m[valid_indices[0]] if len(valid_indices) else data.position_obs_m[0]
    increments = 0.5 * (data.speed_obs_mps[:-1] + data.speed_obs_mps[1:]) * dt_s
    x[1:] = x[0] + np.cumsum(increments)
    if len(valid_indices):
        offset = np.median(data.position_obs_m[valid_indices] - x[valid_indices])
        x += offset
    return x


def line_arrays(
    section: StationSection,
    local_position_m: np.ndarray,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    gradients = np.zeros_like(local_position_m, dtype=float)
    curve_radii = np.full_like(local_position_m, np.inf, dtype=float)
    speed_limits = np.zeros_like(local_position_m, dtype=float)
    for idx, local_pos in enumerate(local_position_m):
        segment = section.segment_at(section.start_m + float(local_pos))
        gradients[idx] = segment.gradient_per_mille
        curve_radii[idx] = np.inf if segment.radius_m is None else segment.radius_m
        speed_limits[idx] = segment.speed_limit_kmh
    return gradients, curve_radii, speed_limits


def resistance_torch(
    speed_mps: torch.Tensor,
    gradient_per_mille: torch.Tensor,
    radius_m: torch.Tensor,
    dynamics: TrainDynamics,
) -> torch.Tensor:
    cfg = dynamics.config
    speed_kmh = speed_mps * 3.6
    davis = cfg.davis_a * speed_kmh**2 + cfg.davis_b * speed_kmh + cfg.davis_c
    gradient = cfg.mass_kg * cfg.gravity_mps2 * gradient_per_mille / 1000.0
    curve_specific = torch.where(
        torch.isfinite(radius_m),
        cfg.curve_resistance_coefficient / torch.clamp(radius_m, min=1.0),
        torch.zeros_like(radius_m),
    )
    curve = cfg.mass_kg * cfg.gravity_mps2 * curve_specific / 1000.0
    return davis + gradient + curve


def residual_metrics(
    speed_mps: np.ndarray,
    position_m: np.ndarray,
    force_n: np.ndarray,
    gradients: np.ndarray,
    radii: np.ndarray,
    dynamics: TrainDynamics,
    dt_s: float,
) -> Dict[str, float]:
    v_mid = 0.5 * (speed_mps[:-1] + speed_mps[1:])
    force_mid = 0.5 * (force_n[:-1] + force_n[1:])
    grad_mid = 0.5 * (gradients[:-1] + gradients[1:])
    radius_mid = np.where(np.isfinite(radii[:-1]), radii[:-1], radii[1:])
    cfg = dynamics.config
    speed_kmh = v_mid * 3.6
    davis = cfg.davis_a * speed_kmh**2 + cfg.davis_b * speed_kmh + cfg.davis_c
    gradient = cfg.mass_kg * cfg.gravity_mps2 * grad_mid / 1000.0
    curve = np.where(
        np.isfinite(radius_mid),
        cfg.mass_kg * cfg.gravity_mps2 * (cfg.curve_resistance_coefficient / radius_mid) / 1000.0,
        0.0,
    )
    dynamics_residual = cfg.mass_kg * np.diff(speed_mps) / dt_s - force_mid + davis + gradient + curve
    kinematic_residual = np.diff(position_m) / dt_s - v_mid
    return {
        "dynamics_rms_n": float(np.sqrt(np.mean(dynamics_residual**2))),
        "dynamics_mean_abs_n": float(np.mean(np.abs(dynamics_residual))),
        "kinematic_rms_mps": float(np.sqrt(np.mean(kinematic_residual**2))),
        "kinematic_mean_abs_mps": float(np.mean(np.abs(kinematic_residual))),
    }


def calibrate_sheet(
    data: SheetCalibrationData,
    section: StationSection,
    base_config: TrainConfig,
    weights: CalibrationWeights,
    iterations: int,
    learning_rate: float,
    dt_s: float,
    device_name: str,
) -> Tuple[Dict[str, np.ndarray], Dict[str, float]]:
    config = replace(base_config, mass_kg=data.mass_kg)
    dynamics = TrainDynamics(config)
    x_base = initial_position_from_speed(data, dt_s)
    gradients, radii, speed_limits = line_arrays(section, x_base)

    device = torch.device(device_name)
    v_obs = torch.as_tensor(data.speed_obs_mps, dtype=torch.float32, device=device)
    x_obs = torch.as_tensor(data.position_obs_m, dtype=torch.float32, device=device)
    x_base_t = torch.as_tensor(x_base, dtype=torch.float32, device=device)
    f_obs = torch.as_tensor(data.force_obs_n, dtype=torch.float32, device=device)
    x_valid = torch.as_tensor(data.position_valid, dtype=torch.bool, device=device)
    grad_t = torch.as_tensor(gradients, dtype=torch.float32, device=device)
    radius_t = torch.as_tensor(radii, dtype=torch.float32, device=device)

    dv = torch.zeros_like(v_obs, requires_grad=True)
    dx = torch.zeros_like(x_base_t, requires_grad=True)
    df = torch.zeros_like(f_obs, requires_grad=True)
    optimizer = torch.optim.Adam([dv, dx, df], lr=learning_rate)

    for _ in range(iterations):
        optimizer.zero_grad()
        v = torch.clamp(v_obs + dv, min=0.0)
        x = x_base_t + dx
        force = f_obs + df
        v_mid = 0.5 * (v[:-1] + v[1:])
        force_mid = 0.5 * (force[:-1] + force[1:])
        grad_mid = 0.5 * (grad_t[:-1] + grad_t[1:])
        radius_mid = torch.where(torch.isfinite(radius_t[:-1]), radius_t[:-1], radius_t[1:])
        resistance = resistance_torch(v_mid, grad_mid, radius_mid, dynamics)
        dynamics_residual = config.mass_kg * (v[1:] - v[:-1]) / dt_s - force_mid + resistance
        kinematic_residual = (x[1:] - x[:-1]) / dt_s - v_mid
        loss = (
            torch.mean(((v - v_obs) / weights.speed_sigma_mps) ** 2)
            + torch.mean(((force - f_obs) / weights.force_sigma_n) ** 2)
            + torch.mean((dynamics_residual / weights.dynamics_sigma_n) ** 2)
            + torch.mean((kinematic_residual / weights.kinematic_sigma_mps) ** 2)
            + 0.15 * torch.mean((torch.diff(force) / weights.force_smooth_sigma_n) ** 2)
            + 0.08 * torch.mean((torch.diff(v) / weights.speed_smooth_sigma_mps) ** 2)
            + 0.50 * torch.mean(torch.relu(-(x[1:] - x[:-1])) / weights.position_step_sigma_m)
        )
        if torch.any(x_valid):
            loss = loss + torch.mean(((x[x_valid] - x_obs[x_valid]) / weights.position_sigma_m) ** 2)
        loss.backward()
        optimizer.step()

    with torch.no_grad():
        v_cal = torch.clamp(v_obs + dv, min=0.0).cpu().numpy()
        x_cal = (x_base_t + dx).cpu().numpy()
        f_cal = (f_obs + df).cpu().numpy()

    before = residual_metrics(
        data.speed_obs_mps,
        x_base,
        data.force_obs_n,
        gradients,
        radii,
        dynamics,
        dt_s,
    )
    after = residual_metrics(v_cal, x_cal, f_cal, gradients, radii, dynamics, dt_s)
    summary = {
        "sheet_name": data.sheet_name,
        "sample_count": len(data.time_s),
        "valid_position_count": int(np.sum(data.position_valid)),
        "mass_kg": data.mass_kg,
        "speed_adjust_rms_kmh": float(np.sqrt(np.mean(((v_cal - data.speed_obs_mps) * 3.6) ** 2))),
        "position_adjust_rms_m": float(np.sqrt(np.mean((x_cal - x_base) ** 2))),
        "force_adjust_rms_kn": float(np.sqrt(np.mean(((f_cal - data.force_obs_n) / 1000.0) ** 2))),
        **{f"before_{key}": value for key, value in before.items()},
        **{f"after_{key}": value for key, value in after.items()},
    }
    trajectory = {
        "time_s": data.time_s,
        "speed_obs_kmh": data.speed_obs_mps * 3.6,
        "speed_calibrated_kmh": v_cal * 3.6,
        "position_base_m": x_base,
        "position_observed_m": data.position_obs_m,
        "position_calibrated_m": x_cal,
        "position_valid": data.position_valid.astype(int),
        "force_obs_kn": data.force_obs_n / 1000.0,
        "force_calibrated_kn": f_cal / 1000.0,
        "gradient_per_mille": gradients,
        "radius_m": radii,
        "speed_limit_kmh": speed_limits,
    }
    return trajectory, summary


def write_csv(path: Path, rows: Sequence[Dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        return
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_trajectory_csv(path: Path, trajectory: Dict[str, np.ndarray]) -> None:
    keys = list(trajectory.keys())
    rows = []
    for idx in range(len(trajectory["time_s"])):
        rows.append({key: trajectory[key][idx] for key in keys})
    write_csv(path, rows)


def plot_speed_comparison(
    trajectories: Sequence[Tuple[str, Dict[str, np.ndarray]]],
    output_dir: Path,
) -> None:
    rows, cols = 4, 3
    fig, axes = plt.subplots(rows, cols, figsize=(16, 12), sharex=False)
    for axis, (sheet_name, trajectory) in zip(axes.flat, trajectories):
        time_s = trajectory["time_s"]
        axis.plot(time_s, trajectory["speed_obs_kmh"], color="#9E9E9E", linewidth=0.8, label="Observed")
        axis.plot(time_s, trajectory["speed_calibrated_kmh"], color="#1565C0", linewidth=0.9, label="Calibrated")
        axis.set_title(sheet_name)
        axis.grid(alpha=0.25)
    for axis in axes.flat[len(trajectories):]:
        axis.set_visible(False)
    axes.flat[0].legend()
    for axis in axes[-1, :]:
        axis.set_xlabel("Time (s)")
    for axis in axes[:, 0]:
        axis.set_ylabel("Speed (km/h)")
    fig.suptitle("THP-TD observed vs physics-calibrated speed")
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    fig.savefig(output_dir / "thp_td_speed_observed_vs_calibrated.png", dpi=180)
    plt.close(fig)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Calibrate THP-TD real data with soft train-dynamics residual constraints."
    )
    parser.add_argument("--input", type=Path, default=default_thp_td_path())
    parser.add_argument("--line-dir", type=Path, default=DEFAULT_LINE_DIR)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent / "thp_td_dynamics_calibration")
    parser.add_argument("--sheet-start", type=int, default=2)
    parser.add_argument("--sheet-end", type=int, default=13)
    parser.add_argument("--dt", type=float, default=0.1)
    parser.add_argument("--iterations", type=int, default=700)
    parser.add_argument("--learning-rate", type=float, default=0.035)
    parser.add_argument("--device", choices=("cpu", "cuda"), default="cpu")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    profile = LineProfile.from_excel(args.line_dir / "slope.xlsx", args.line_dir / "curvature.xlsx")
    section = profile.section_by_name["THP-TD"]
    base_config = TrainConfig(name="Real THP-TD")
    weights = CalibrationWeights()
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    sheet_names = [
        f"Sheet{idx}"
        for idx in range(args.sheet_start, args.sheet_end + 1)
        if f"Sheet{idx}" in workbook.sheetnames
    ]
    summaries: List[Dict[str, object]] = []
    plot_items: List[Tuple[str, Dict[str, np.ndarray]]] = []
    for sheet_name in sheet_names:
        print(f"Calibrating {sheet_name}")
        data = read_sheet_data(args.input, sheet_name, args.dt, base_config.mass_kg)
        trajectory, summary = calibrate_sheet(
            data,
            section,
            base_config,
            weights,
            args.iterations,
            args.learning_rate,
            args.dt,
            args.device,
        )
        summaries.append(summary)
        plot_items.append((sheet_name, trajectory))
        write_trajectory_csv(args.output_dir / f"{sheet_name}_calibrated_trajectory.csv", trajectory)
        print(
            f"  dynamics RMS: {summary['before_dynamics_rms_n'] / 1000.0:.1f} -> "
            f"{summary['after_dynamics_rms_n'] / 1000.0:.1f} kN, "
            f"speed adjust RMS={summary['speed_adjust_rms_kmh']:.3f} km/h"
        )
    write_csv(args.output_dir / "thp_td_calibration_summary.csv", summaries)
    plot_speed_comparison(plot_items, args.output_dir)
    print(f"Calibration reports written to: {args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
