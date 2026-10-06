"""Read and merge slope, curvature and inherited speed-limit information."""

from __future__ import annotations

import bisect
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

try:
    from .config import EPS, INITIAL_SPEED_LIMIT_KMH, SECTION_NAMES
except ImportError:
    from config import EPS, INITIAL_SPEED_LIMIT_KMH, SECTION_NAMES


@dataclass(frozen=True)
class SlopeSegment:
    start_m: float
    end_m: float
    gradient_per_mille: float


@dataclass(frozen=True)
class CurveSegment:
    start_m: float
    end_m: float
    radius_m: float
    speed_limit_kmh: float


@dataclass(frozen=True)
class TrackSegment:
    start_m: float
    end_m: float
    gradient_per_mille: float
    radius_m: Optional[float]
    speed_limit_kmh: float
    physical_start_m: Optional[float] = None
    physical_end_m: Optional[float] = None


@dataclass
class StationSection:
    name: str
    start_m: float
    end_m: float
    segments: List[TrackSegment]
    physical_start_m: Optional[float] = None
    physical_end_m: Optional[float] = None

    def __post_init__(self) -> None:
        self._segment_starts = [item.start_m for item in self.segments]

    @property
    def length_m(self) -> float:
        return self.end_m - self.start_m

    @property
    def display_start_m(self) -> float:
        return self.start_m if self.physical_start_m is None else self.physical_start_m

    @property
    def display_end_m(self) -> float:
        return self.end_m if self.physical_end_m is None else self.physical_end_m

    def physical_position_m(self, position_m: float) -> float:
        if self.physical_start_m is None or self.physical_end_m is None:
            return position_m
        ratio = (position_m - self.start_m) / max(self.length_m, EPS)
        return self.physical_start_m + ratio * (self.physical_end_m - self.physical_start_m)

    def segment_at(self, position_m: float) -> TrackSegment:
        clamped = min(max(position_m, self.start_m), self.end_m - EPS)
        idx = bisect.bisect_right(self._segment_starts, clamped) - 1
        return self.segments[max(0, idx)]

    def future_segments(self, position_m: float) -> Iterable[TrackSegment]:
        idx = bisect.bisect_right(self._segment_starts, position_m)
        return self.segments[idx:]


class LineProfile:
    def __init__(self, sections: List[StationSection]) -> None:
        self.sections = sections
        self.section_by_name = {item.name: item for item in sections}

    @classmethod
    def from_excel(cls, slope_path: Path, curvature_path: Path) -> "LineProfile":
        slope_wb = load_workbook(slope_path, read_only=True, data_only=True)
        curve_wb = load_workbook(curvature_path, read_only=True, data_only=True)
        if slope_wb.sheetnames != curve_wb.sheetnames:
            raise ValueError("slope.xlsx and curvature.xlsx sheet orders do not match")
        if slope_wb.sheetnames != SECTION_NAMES:
            raise ValueError(f"Unexpected section order: {slope_wb.sheetnames}")

        slopes_by_section: Dict[str, List[SlopeSegment]] = {}
        curves_by_section: Dict[str, List[CurveSegment]] = {}
        speed_events: List[Tuple[float, float]] = [(346.0, INITIAL_SPEED_LIMIT_KMH)]
        for name in SECTION_NAMES:
            slopes = cls._read_slopes(name, slope_wb[name])
            curves = cls._read_curves(curve_wb[name])
            cls._validate_source_segments(name, slopes, curves)
            slopes_by_section[name] = slopes
            curves_by_section[name] = curves
            speed_events.extend((item.start_m, item.speed_limit_kmh) for item in curves)

        speed_events.sort()
        sections = [
            cls._merge_section(
                name,
                slopes_by_section[name],
                curves_by_section[name],
                speed_events,
            )
            for name in SECTION_NAMES
        ]
        return cls(sections)

    @classmethod
    def reversed_with_curvature(
        cls,
        up_profile: "LineProfile",
        curvature_path: Path,
    ) -> "LineProfile":
        """Build the down direction using dedicated down curvature and limits."""
        line_end_m = up_profile.sections[-1].end_m
        transformed_curves, speed_events, corrections = (
            cls._read_transformed_down_curves(curvature_path, line_end_m)
        )

        base_down = up_profile.reversed_direction()
        sections = [
            cls._merge_reversed_section(
                section,
                transformed_curves,
                speed_events,
            )
            for section in base_down.sections
        ]
        profile = cls(sections)
        profile.source_corrections = tuple(corrections)
        return profile

    @classmethod
    def reversed_with_slope_and_curvature(
        cls,
        up_profile: "LineProfile",
        slope_path: Path,
        curvature_path: Path,
    ) -> "LineProfile":
        """Build the down direction using dedicated down slope, curvature and limits."""
        line_end_m = up_profile.sections[-1].end_m
        transformed_curves, speed_events, curve_corrections = (
            cls._read_transformed_down_curves(curvature_path, line_end_m)
        )
        transformed_slopes, slope_corrections = cls._read_transformed_down_slopes(
            slope_path,
            up_profile.sections[0].start_m,
            line_end_m,
        )

        base_down = up_profile.reversed_direction()
        sections = [
            cls._merge_down_section(
                section,
                transformed_slopes,
                transformed_curves,
                speed_events,
            )
            for section in base_down.sections
        ]
        profile = cls(sections)
        profile.source_corrections = tuple(curve_corrections + slope_corrections)
        return profile

    @classmethod
    def _read_transformed_down_curves(
        cls,
        curvature_path: Path,
        line_end_m: float,
    ) -> Tuple[List[CurveSegment], List[Tuple[float, float]], List[str]]:
        curve_wb = load_workbook(curvature_path, read_only=True, data_only=True)
        if curve_wb.sheetnames != SECTION_NAMES:
            raise ValueError(
                f"Unexpected down-curvature sheet order: {curve_wb.sheetnames}"
            )

        down_curves: List[CurveSegment] = []
        corrections: List[str] = []
        for name in SECTION_NAMES:
            worksheet = curve_wb[name]
            for row_idx, values in enumerate(
                worksheet.iter_rows(
                    min_row=2,
                    min_col=1,
                    max_col=11,
                    values_only=True,
                ),
                start=2,
            ):
                if all(value is None for value in values):
                    continue
                direction = str(values[1]).strip()
                if direction != "\u4e0b\u884c":
                    raise ValueError(
                        f"Expected down-direction data in {name} row {row_idx}, "
                        f"got {values[1]!r}"
                    )
                start_m = cls._parse_mileage(
                    values[2],
                    line_end_m,
                    name,
                    row_idx,
                    "start",
                    corrections,
                )
                end_m = cls._parse_mileage(
                    values[3],
                    line_end_m,
                    name,
                    row_idx,
                    "end",
                    corrections,
                )
                radius_m = values[4]
                speed_limit = values[10]
                if radius_m is None or speed_limit is None:
                    raise ValueError(
                        f"Incomplete down-curvature row in {name} row {row_idx}"
                    )
                if end_m <= start_m:
                    raise ValueError(
                        f"Invalid down-curvature interval in {name} row {row_idx}: "
                        f"{start_m} -> {end_m}"
                    )
                down_curves.append(
                    CurveSegment(
                        start_m,
                        end_m,
                        float(radius_m),
                        float(speed_limit),
                    )
                )

        down_curves.sort(key=lambda item: item.start_m)
        for previous, current in zip(down_curves, down_curves[1:]):
            if current.start_m < previous.end_m - EPS:
                raise ValueError(
                    "Overlapping down-curvature intervals: "
                    f"{previous} and {current}"
                )

        transformed_curves = [
            CurveSegment(
                line_end_m - curve.end_m,
                line_end_m - curve.start_m,
                curve.radius_m,
                curve.speed_limit_kmh,
            )
            for curve in reversed(down_curves)
        ]
        speed_events = [
            (curve.start_m, curve.speed_limit_kmh)
            for curve in transformed_curves
        ]
        if not speed_events:
            raise ValueError("No down-direction curvature data found")
        speed_events.insert(0, (0.0, speed_events[0][1]))
        return transformed_curves, speed_events, corrections

    @classmethod
    def _read_transformed_down_slopes(
        cls,
        slope_path: Path,
        line_start_m: float,
        line_end_m: float,
    ) -> Tuple[List[SlopeSegment], List[str]]:
        slope_wb = load_workbook(slope_path, read_only=True, data_only=True)
        worksheet = slope_wb[slope_wb.sheetnames[0]]
        physical_slopes: List[SlopeSegment] = []
        corrections: List[str] = []
        for row_idx, values in enumerate(
            worksheet.iter_rows(
                min_row=2,
                min_col=1,
                max_col=5,
                values_only=True,
            ),
            start=2,
        ):
            if all(value is None for value in values):
                continue
            direction = str(values[0]).strip()
            if direction not in {"\u4e0b\u884c", "\u00cf\u00c2\u00d0\u00d0"}:
                continue
            start_m = cls._parse_mileage(
                values[1],
                line_end_m,
                worksheet.title,
                row_idx,
                "slope_start",
                corrections,
                allow_outside=True,
            )
            if values[3] is None:
                raise ValueError(f"Missing down-slope length in row {row_idx}")
            end_m = start_m + float(values[3])
            gradient = values[4]
            if gradient is None:
                raise ValueError(f"Missing down-slope gradient in row {row_idx}")
            if end_m <= start_m:
                corrections.append(
                    f"{worksheet.title} row {row_idx} slope interval reversed: "
                    f"{start_m} -> {end_m}"
                )
                start_m, end_m = end_m, start_m
            clipped_start = max(line_start_m, start_m)
            clipped_end = min(line_end_m, end_m)
            if clipped_end <= clipped_start:
                continue
            physical_slopes.append(
                SlopeSegment(
                    clipped_start,
                    clipped_end,
                    -float(gradient),
                )
            )

        if not physical_slopes:
            raise ValueError(f"No usable down-slope rows found in {slope_path}")
        physical_slopes.sort(key=lambda item: item.start_m)
        continuous_slopes: List[SlopeSegment] = []
        for slope in physical_slopes:
            start_m = max(line_start_m, slope.start_m)
            end_m = min(line_end_m, slope.end_m)
            if end_m <= line_start_m or start_m >= line_end_m:
                continue
            if continuous_slopes:
                previous = continuous_slopes[-1]
                if start_m > previous.end_m + EPS:
                    corrections.append(
                        "Down-slope length ended before next change mileage; "
                        f"extended {previous.end_m} -> {start_m}"
                    )
                    continuous_slopes[-1] = SlopeSegment(
                        previous.start_m,
                        start_m,
                        previous.gradient_per_mille,
                    )
                elif start_m < previous.end_m - EPS:
                    corrections.append(
                        "Down-slope length passed next change mileage; "
                        f"truncated {previous.end_m} -> {start_m}"
                    )
                    if start_m <= previous.start_m + EPS:
                        continuous_slopes.pop()
                    else:
                        continuous_slopes[-1] = SlopeSegment(
                            previous.start_m,
                            start_m,
                            previous.gradient_per_mille,
                        )
            if end_m <= start_m + EPS:
                continue
            continuous_slopes.append(
                SlopeSegment(start_m, end_m, slope.gradient_per_mille)
            )

        if not continuous_slopes:
            raise ValueError(f"No clipped down-slope rows remain in {slope_path}")
        if continuous_slopes[0].start_m > line_start_m + EPS:
            corrections.append(
                "Down-slope starts after line start; extended first segment "
                f"{line_start_m} -> {continuous_slopes[0].start_m}"
            )
            first = continuous_slopes[0]
            continuous_slopes[0] = SlopeSegment(
                line_start_m,
                first.end_m,
                first.gradient_per_mille,
            )
        if continuous_slopes[-1].end_m < line_end_m - EPS:
            corrections.append(
                "Down-slope ends before line end; extended last segment "
                f"{continuous_slopes[-1].end_m} -> {line_end_m}"
            )
            last = continuous_slopes[-1]
            continuous_slopes[-1] = SlopeSegment(
                last.start_m,
                line_end_m,
                last.gradient_per_mille,
            )

        for previous, current in zip(continuous_slopes, continuous_slopes[1:]):
            if current.start_m > previous.end_m + EPS:
                corrections.append(
                    "Down-slope gap after clipping: "
                    f"{previous.end_m} -> {current.start_m}"
                )
            elif current.start_m < previous.end_m - EPS:
                corrections.append(
                    "Down-slope overlap after clipping: "
                    f"{current.start_m} -> {previous.end_m}"
                )

        transformed_slopes = [
            SlopeSegment(
                line_end_m - slope.end_m,
                line_end_m - slope.start_m,
                slope.gradient_per_mille,
            )
            for slope in reversed(continuous_slopes)
        ]
        return transformed_slopes, corrections

    @staticmethod
    def _parse_mileage(
        value,
        line_end_m: float,
        section_name: str,
        row_idx: int,
        field_name: str,
        corrections: List[str],
        allow_outside: bool = False,
    ) -> float:
        if isinstance(value, str):
            text = value.strip().upper()
            match = re.fullmatch(r"K(\d+)\+(\d+(?:\.\d+)?)", text)
            if match:
                parsed = float(match.group(1)) * 1000.0 + float(match.group(2))
                corrections.append(
                    f"{section_name} row {row_idx} {field_name}: {value!r} -> {parsed}"
                )
            else:
                match = re.fullmatch(r"K(\d+(?:\.\d+)?)", text)
                if match:
                    parsed = float(match.group(1))
                    corrections.append(
                        f"{section_name} row {row_idx} {field_name}: "
                        f"{value!r} -> {parsed}"
                    )
                else:
                    raise ValueError(
                        f"Unsupported mileage {value!r} in {section_name} row {row_idx}"
                    )
        elif value is None:
            raise ValueError(
                f"Missing mileage in {section_name} row {row_idx} ({field_name})"
            )
        else:
            parsed = float(value)
        original = parsed
        while parsed > line_end_m + 1000.0:
            parsed /= 1000.0
        if not math.isclose(parsed, original):
            corrections.append(
                f"{section_name} row {row_idx} {field_name}: {original} -> {parsed}"
            )
        if not allow_outside and not 0.0 <= parsed <= line_end_m:
            raise ValueError(
                f"Mileage outside line range in {section_name} row {row_idx}: "
                f"{parsed}"
            )
        return parsed

    @staticmethod
    def _merge_reversed_section(
        base_section: StationSection,
        curves: Sequence[CurveSegment],
        speed_events: Sequence[Tuple[float, float]],
    ) -> StationSection:
        breakpoints = {base_section.start_m, base_section.end_m}
        for segment in base_section.segments:
            breakpoints.update((segment.start_m, segment.end_m))
        for curve in curves:
            if curve.end_m <= base_section.start_m or curve.start_m >= base_section.end_m:
                continue
            breakpoints.add(max(base_section.start_m, curve.start_m))
            breakpoints.add(min(base_section.end_m, curve.end_m))
        event_positions = [item[0] for item in speed_events]
        merged: List[TrackSegment] = []
        points = sorted(breakpoints)
        for start_m, end_m in zip(points, points[1:]):
            midpoint = (start_m + end_m) / 2.0
            slope_segment = base_section.segment_at(midpoint)
            curve = next(
                (
                    item
                    for item in curves
                    if item.start_m <= midpoint < item.end_m
                ),
                None,
            )
            event_idx = bisect.bisect_right(event_positions, start_m) - 1
            merged.append(
                TrackSegment(
                    start_m=start_m,
                    end_m=end_m,
                    gradient_per_mille=slope_segment.gradient_per_mille,
                    radius_m=None if curve is None else curve.radius_m,
                    speed_limit_kmh=speed_events[max(0, event_idx)][1],
                    physical_start_m=base_section.physical_position_m(start_m),
                    physical_end_m=base_section.physical_position_m(end_m),
                )
            )
        return StationSection(
            base_section.name,
            base_section.start_m,
            base_section.end_m,
            merged,
            base_section.physical_start_m,
            base_section.physical_end_m,
        )

    @staticmethod
    def _merge_down_section(
        base_section: StationSection,
        slopes: Sequence[SlopeSegment],
        curves: Sequence[CurveSegment],
        speed_events: Sequence[Tuple[float, float]],
    ) -> StationSection:
        breakpoints = {base_section.start_m, base_section.end_m}
        for slope in slopes:
            if slope.end_m <= base_section.start_m or slope.start_m >= base_section.end_m:
                continue
            breakpoints.add(max(base_section.start_m, slope.start_m))
            breakpoints.add(min(base_section.end_m, slope.end_m))
        for curve in curves:
            if curve.end_m <= base_section.start_m or curve.start_m >= base_section.end_m:
                continue
            breakpoints.add(max(base_section.start_m, curve.start_m))
            breakpoints.add(min(base_section.end_m, curve.end_m))

        event_positions = [item[0] for item in speed_events]
        merged: List[TrackSegment] = []
        points = sorted(breakpoints)
        for start_m, end_m in zip(points, points[1:]):
            midpoint = (start_m + end_m) / 2.0
            slope = next(
                (
                    item
                    for item in slopes
                    if item.start_m <= midpoint < item.end_m
                ),
                None,
            )
            if slope is None:
                raise ValueError(
                    f"No down-slope coverage in {base_section.name} at "
                    f"{midpoint:.3f} m"
                )
            curve = next(
                (
                    item
                    for item in curves
                    if item.start_m <= midpoint < item.end_m
                ),
                None,
            )
            event_idx = bisect.bisect_right(event_positions, start_m) - 1
            merged.append(
                TrackSegment(
                    start_m=start_m,
                    end_m=end_m,
                    gradient_per_mille=slope.gradient_per_mille,
                    radius_m=None if curve is None else curve.radius_m,
                    speed_limit_kmh=speed_events[max(0, event_idx)][1],
                    physical_start_m=base_section.physical_position_m(start_m),
                    physical_end_m=base_section.physical_position_m(end_m),
                )
            )
        return StationSection(
            base_section.name,
            base_section.start_m,
            base_section.end_m,
            merged,
            base_section.physical_start_m,
            base_section.physical_end_m,
        )

    @staticmethod
    def _read_slopes(name: str, worksheet) -> List[SlopeSegment]:
        rows: List[SlopeSegment] = []
        for start_m, gradient, length_m in worksheet.iter_rows(
            min_row=2,
            min_col=1,
            max_col=3,
            values_only=True,
        ):
            if start_m is None and gradient is None and length_m is None:
                continue
            if start_m is None or gradient is None or length_m is None:
                raise ValueError(f"Incomplete slope row in {name}")
            start = float(start_m)
            length = float(length_m)
            if name == "SC-SY" and math.isclose(start, 15850.0):
                length = 600.0
            rows.append(SlopeSegment(start, start + length, float(gradient)))
        return rows

    @staticmethod
    def _read_curves(worksheet) -> List[CurveSegment]:
        rows: List[CurveSegment] = []
        for values in worksheet.iter_rows(min_row=2, min_col=1, max_col=11, values_only=True):
            if all(value is None for value in values):
                continue
            start_m, end_m, radius_m, speed_limit = values[2], values[3], values[4], values[10]
            if None in (start_m, end_m, radius_m, speed_limit):
                raise ValueError(f"Incomplete curvature row in {worksheet.title}")
            rows.append(
                CurveSegment(
                    float(start_m),
                    float(end_m),
                    float(radius_m),
                    float(speed_limit),
                )
            )
        return rows

    @staticmethod
    def _validate_source_segments(
        name: str,
        slopes: Sequence[SlopeSegment],
        curves: Sequence[CurveSegment],
    ) -> None:
        if not slopes:
            raise ValueError(f"No slope data in {name}")
        for previous, current in zip(slopes, slopes[1:]):
            if not math.isclose(previous.end_m, current.start_m, abs_tol=1e-6):
                raise ValueError(
                    f"Slope segments are not continuous in {name}: "
                    f"{previous.end_m} -> {current.start_m}"
                )

        section_start = slopes[0].start_m
        section_end = slopes[-1].end_m
        previous_end = section_start
        for curve in curves:
            if curve.start_m < section_start or curve.end_m > section_end:
                raise ValueError(f"Curve outside station section in {name}: {curve}")
            if curve.start_m < previous_end or curve.end_m <= curve.start_m:
                raise ValueError(f"Invalid or overlapping curve in {name}: {curve}")
            previous_end = curve.end_m

    @staticmethod
    def _merge_section(
        name: str,
        slopes: Sequence[SlopeSegment],
        curves: Sequence[CurveSegment],
        speed_events: Sequence[Tuple[float, float]],
    ) -> StationSection:
        section_start = slopes[0].start_m
        section_end = slopes[-1].end_m
        breakpoints = {section_start, section_end}
        for slope in slopes:
            breakpoints.update((slope.start_m, slope.end_m))
        for curve in curves:
            breakpoints.update((curve.start_m, curve.end_m))

        event_positions = [item[0] for item in speed_events]
        merged: List[TrackSegment] = []
        points = sorted(breakpoints)
        for start_m, end_m in zip(points, points[1:]):
            midpoint = (start_m + end_m) / 2.0
            slope = next(item for item in slopes if item.start_m <= midpoint < item.end_m)
            curve = next(
                (item for item in curves if item.start_m <= midpoint < item.end_m),
                None,
            )
            event_idx = bisect.bisect_right(event_positions, start_m) - 1
            merged.append(
                TrackSegment(
                    start_m=start_m,
                    end_m=end_m,
                    gradient_per_mille=slope.gradient_per_mille,
                    radius_m=None if curve is None else curve.radius_m,
                    speed_limit_kmh=speed_events[max(0, event_idx)][1],
                )
            )
        return StationSection(name, section_start, section_end, merged)

    def print_summary(self) -> None:
        print("Loaded line profile")
        print(f"  sections: {len(self.sections)}")
        print(f"  line range: {self.sections[0].start_m:.3f} .. {self.sections[-1].end_m:.3f} m")
        print(f"  merged track segments: {sum(len(item.segments) for item in self.sections)}")
        print(f"  initial speed limit: {INITIAL_SPEED_LIMIT_KMH:.1f} km/h")
        print("  SC-SY slope at 15850 m: patched to 600 m")

    def reversed_direction(self) -> "LineProfile":
        """Build the opposite direction using the same physical alignment."""
        line_end_m = self.sections[-1].end_m
        reversed_sections: List[StationSection] = []
        for section in reversed(self.sections):
            reversed_segments = [
                TrackSegment(
                    start_m=line_end_m - segment.end_m,
                    end_m=line_end_m - segment.start_m,
                    gradient_per_mille=-segment.gradient_per_mille,
                    radius_m=segment.radius_m,
                    speed_limit_kmh=segment.speed_limit_kmh,
                    physical_start_m=segment.end_m,
                    physical_end_m=segment.start_m,
                )
                for segment in reversed(section.segments)
            ]
            start_station, end_station = section.name.split("-", maxsplit=1)
            reversed_sections.append(
                StationSection(
                    name=f"{end_station}-{start_station}",
                    start_m=line_end_m - section.end_m,
                    end_m=line_end_m - section.start_m,
                    segments=reversed_segments,
                    physical_start_m=section.end_m,
                    physical_end_m=section.start_m,
                )
            )
        return LineProfile(reversed_sections)
